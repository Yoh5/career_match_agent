"""Exports bureautiques : lettre de motivation en .docx, CV en .pdf.

- Lettre → Word via python-docx (déjà utilisé pour LIRE les CV) : marges A4,
  interligne aéré, prête à personnaliser/envoyer.
- CV → PDF via fpdf2 (pur Python, zéro dépendance système — là où career-ops
  passe par Playwright/Chromium).

**Deux mises en page, pour deux lecteurs différents :**

| `layout` | Pour qui | Forme |
|---|---|---|
| `"ats"` (défaut) | le robot de tri | **une seule colonne**, sections standard, noir sur blanc, puces `-` |
| `"designed"` | l'œil humain | 2 colonnes, bandeau couleur, barre latérale compétences |

C'est une distinction qui compte : un PDF à deux colonnes est extrait par les
parseurs ATS dans l'ordre du flux de texte, ce qui **entrelace la barre latérale
avec le corps du CV** et peut coller « Docker » au milieu d'une phrase sur un
stage. Tant que le CV passe par un ATS, la version une colonne est celle qu'il
faut envoyer — `atscheck.py` mesure précisément ce que chaque version perd.

Tout retourne des bytes ; aucune écriture disque, aucune dépendance réseau.
"""
import io
import re
from typing import Dict, List, Optional

# ── Couleurs de la charte (identiques au template HTML) ────────────────────
_ACCENT = (43, 108, 176)     # bleu
_INK = (26, 32, 44)
_MUTED = (74, 85, 104)
_LINE = (226, 232, 240)

_LABELS = {
    "fr": {"profile": "Profil", "experience": "Expériences", "projects": "Projets",
           "education": "Formation", "skills": "Compétences", "certs": "Certifications",
           "langs": "Langues", "links": "Liens"},
    "en": {"profile": "Profile", "experience": "Experience", "projects": "Projects",
           "education": "Education", "skills": "Skills", "certs": "Certifications",
           "langs": "Languages", "links": "Links"},
}

# Intitulés de section de la mise en page ATS : les libellés CANONIQUES que les
# parseurs cherchent pour découper un CV. « Expérience professionnelle » est
# reconnu, « Mon parcours » ne l'est pas — d'où des titres volontairement plats.
_ATS_LABELS = {
    "fr": {"profile": "PROFIL", "experience": "EXPÉRIENCE PROFESSIONNELLE",
           "projects": "PROJETS", "education": "FORMATION", "skills": "COMPÉTENCES",
           "certs": "CERTIFICATIONS", "langs": "LANGUES"},
    "en": {"profile": "PROFESSIONAL SUMMARY", "experience": "PROFESSIONAL EXPERIENCE",
           "projects": "PROJECTS", "education": "EDUCATION", "skills": "SKILLS",
           "certs": "CERTIFICATIONS", "langs": "LANGUAGES"},
}


def _latin(s: str) -> str:
    """fpdf2 en polices de base = latin-1 : on remplace les caractères hors charte.

    Tout ce qui n'est pas traduit ici finit en « ? » dans le PDF — un « CV ↔ offre »
    devient « CV ? offre » et le CV part avec une faute sous les yeux du recruteur.
    La table couvre donc les symboles qu'on croise vraiment dans un CV tech."""
    repl = {"—": "-", "–": "-", "‘": "'", "’": "'", "“": '"',
            "”": '"', "•": "-", "…": "...", " ": " ", "→": "->",
            "·": "-", "✅": "", "▸": "-", "↔": "<->", "←": "<-",
            "⇒": "=>", "≥": ">=", "≤": "<=", "×": "x", "≠": "!=",
            "€": "EUR", "™": "(TM)", "✓": "", "★": "", "▪": "-"}
    for a, b in repl.items():
        s = (s or "").replace(a, b)
    return s.encode("latin-1", "replace").decode("latin-1")


# ═══ Lettre de motivation → .docx ═══════════════════════════════════════════

def letter_docx(text: str) -> bytes:
    """Lettre en Word : paragraphes du texte, marges A4, police lisible."""
    from docx import Document
    from docx.shared import Pt, Cm
    from docx.enum.text import WD_LINE_SPACING

    doc = Document()
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Cm(21), Cm(29.7)
    sec.top_margin = sec.bottom_margin = Cm(2.2)
    sec.left_margin = sec.right_margin = Cm(2.4)

    style = doc.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(11)

    for block in re.split(r"\n\s*\n", (text or "").strip()):
        p = doc.add_paragraph(block.strip())
        p.paragraph_format.space_after = Pt(10)
        p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.MULTIPLE
        p.paragraph_format.line_spacing = 1.15

    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


# ═══ CV → .pdf (2 colonnes, depuis le JSON structuré) ═══════════════════════

_M = 12          # marge (mm)
_PAGE_W = 210
_PAGE_H = 297
_MAIN_X, _MAIN_W = _M, 120                  # colonne principale
_SIDE_X = _MAIN_X + _MAIN_W + 8             # barre latérale
_SIDE_W = _PAGE_W - _SIDE_X - _M
_LIMIT_Y = _PAGE_H - _M


class _Col:
    """Curseur de colonne indépendant (x fixe, y qui coule, saut de page géré)."""

    def __init__(self, pdf, x, w, y):
        self.pdf, self.x, self.w, self.y = pdf, x, w, y
        self.page = pdf.page

    def _ensure(self, h):
        if self.y + h > _LIMIT_Y:
            if self.page >= self.pdf.pages_count:
                self.pdf.add_page()
            self.page += 1
            self.pdf.page = self.page
            self.y = _M

    def text(self, s, size=9.2, style="", color=_INK, lh=1.32, before=0.0):
        s = _latin(s).strip()
        if not s:
            return
        self.pdf.page = self.page
        self.pdf.set_font("Helvetica", style, size)
        self.pdf.set_text_color(*color)
        line_h = size * 0.3528 * lh
        self._ensure(line_h + before)
        self.y += before
        self.pdf.set_xy(self.x, self.y)
        self.pdf.multi_cell(self.w, line_h, s)
        self.y = self.pdf.get_y()

    def bullet(self, s, size=9.2):
        s = _latin(s).strip()
        if not s:
            return
        self.pdf.page = self.page
        self.pdf.set_font("Helvetica", "", size)
        line_h = size * 0.3528 * 1.32
        self._ensure(line_h)
        self.pdf.set_text_color(*_ACCENT)
        self.pdf.set_xy(self.x, self.y)
        self.pdf.cell(4, line_h, "-")
        self.pdf.set_text_color(*_INK)
        self.pdf.set_xy(self.x + 4, self.y)
        self.pdf.multi_cell(self.w - 4, line_h, s)
        self.y = self.pdf.get_y()

    def heading(self, s):
        self.pdf.page = self.page
        self._ensure(9)
        self.y += 3.5
        self.text(s.upper(), size=9.6, style="B", color=_ACCENT, before=0)
        self.pdf.page = self.page
        self.pdf.set_draw_color(*_LINE)
        self.pdf.line(self.x, self.y + 0.6, self.x + self.w, self.y + 0.6)
        self.y += 2.4


def _clean(items) -> List:
    return [x for x in (items or []) if x]


def cv_pdf(structured: Optional[Dict], cv_markdown: str = "", lang: str = "fr",
           layout: str = "ats") -> bytes:
    """CV PDF A4.

    - `layout="ats"` (défaut) → **une colonne**, sections standard, monochrome :
      la version à envoyer dès qu'un robot lit le CV.
    - `layout="designed"` → 2 colonnes (bandeau, profil/expériences/projets à
      gauche ; compétences/formation/certifs/langues à droite) : jolie, mais son
      ordre de lecture est illisible pour une partie des parseurs.

    Sans `structured`, les deux retombent sur le rendu Markdown une colonne."""
    from fpdf import FPDF

    lg = "en" if str(lang).lower().startswith("en") else "fr"
    L = _LABELS[lg]
    pdf = FPDF(orientation="P", unit="mm", format="A4")
    pdf.set_auto_page_break(False)
    pdf.set_margins(_M, _M, _M)
    pdf.add_page()

    if not isinstance(structured, dict) or not structured.get("name"):
        return _cv_pdf_fallback(pdf, cv_markdown)

    if str(layout).lower() != "designed":
        return _ats_autofit(structured, lg, first=pdf)

    # ── Bandeau d'en-tête ──
    contact = structured.get("contact") or {}
    contact_bits = _clean([contact.get(k) for k in ("email", "phone", "location", "linkedin", "github", "portfolio")])
    head_h = 30 if contact_bits else 24
    pdf.set_fill_color(*_ACCENT)
    pdf.rect(0, 0, _PAGE_W, head_h, "F")
    pdf.set_text_color(255, 255, 255)
    pdf.set_font("Helvetica", "B", 19)
    pdf.set_xy(_M, 7)
    pdf.cell(0, 8, _latin(structured.get("name") or ""))
    pdf.set_font("Helvetica", "", 11)
    pdf.set_xy(_M, 15)
    pdf.cell(0, 6, _latin(structured.get("role") or ""))
    if contact_bits:
        pdf.set_font("Helvetica", "", 8)
        pdf.set_xy(_M, 22)
        pdf.multi_cell(_PAGE_W - 2 * _M, 4, _latin("  |  ".join(contact_bits)))

    y0 = head_h + 5
    main = _Col(pdf, _MAIN_X, _MAIN_W, y0)
    side = _Col(pdf, _SIDE_X, _SIDE_W, y0)

    # ── Colonne principale ──
    if structured.get("summary"):
        main.heading(L["profile"])
        main.text(structured["summary"])
    for key, label in (("experiences", L["experience"]), ("projects", L["projects"])):
        entries = _clean(structured.get(key))
        if not entries:
            continue
        main.heading(label)
        for e in entries:
            if not isinstance(e, dict):
                main.bullet(str(e))
                continue
            title = e.get("title") or ""
            org = e.get("org") or e.get("meta") or ""
            head = f"{title} - {org}" if org else title
            main.text(head, style="B", before=1.6)
            sub = " | ".join(_clean([e.get("date"), e.get("stack")]))
            if sub:
                main.text(sub, size=8.2, color=_MUTED)
            for b in _clean(e.get("bullets"))[:6]:
                main.bullet(str(b))

    # ── Barre latérale ──
    pdf.page = side.page
    skills = _clean(structured.get("skills"))
    if skills:
        side.heading(L["skills"])
        for g in skills:
            if isinstance(g, dict):
                if g.get("group"):
                    side.text(str(g["group"]), size=8.6, style="B", before=1.2)
                side.text(", ".join(str(i) for i in _clean(g.get("items"))), size=8.4, color=_MUTED)
            else:
                side.bullet(str(g), size=8.4)
    edu = _clean(structured.get("education"))
    if edu:
        side.heading(L["education"])
        for e in edu:
            if isinstance(e, dict):
                side.text(str(e.get("title") or ""), size=8.6, style="B", before=1.2)
                if e.get("meta"):
                    side.text(str(e["meta"]), size=8.2, color=_MUTED)
            else:
                side.text(str(e), size=8.6, before=1.2)
    for key, label in (("certifications", L["certs"]), ("languages", L["langs"])):
        entries = _clean(structured.get(key))
        if entries:
            side.heading(label)
            for e in entries:
                side.bullet(str(e), size=8.4)
    links = _clean([contact.get(k) for k in ("linkedin", "github", "portfolio")])
    if links:
        side.heading(L["links"])
        for u in links:
            side.text(str(u), size=8.0, color=_ACCENT)

    return bytes(pdf.output())


def _ats_autofit(s: Dict, lg: str, first=None) -> bytes:
    """Rend le CV ATS et tente de le faire tenir sur UNE page.

    Un CV junior qui déborde de quelques lignes se lit mal — mais on ne compresse
    pas à l'aveugle : on essaie des paliers de plus en plus serrés et on s'arrête
    au premier qui tient. Si même le plus serré déborde, c'est que le contenu est
    réellement trop long, et on rend la version LISIBLE (échelle 1) sur deux pages
    plutôt qu'un mur de texte illisible."""
    from fpdf import FPDF

    def render(scale, pdf=None):
        if pdf is None:
            pdf = FPDF(orientation="P", unit="mm", format="A4")
            pdf.set_auto_page_break(False)
            pdf.set_margins(_M, _M, _M)
            pdf.add_page()
        return _cv_pdf_ats(pdf, s, lg, scale)

    full = render(1.0, first)
    if _page_count(full) <= 1:
        return full
    for scale in (0.94, 0.88, 0.82):
        data = render(scale)
        if _page_count(data) <= 1:
            return data
    return full


def _page_count(pdf_bytes: bytes) -> int:
    try:
        import io as _io
        from pypdf import PdfReader
        return len(PdfReader(_io.BytesIO(pdf_bytes)).pages)
    except Exception:                       # pragma: no cover - garde-fou
        return 1


def _cv_pdf_ats(pdf, s: Dict, lg: str, scale: float = 1.0) -> bytes:
    """Mise en page UNE COLONNE, pensée pour être ré-extraite proprement.

    Les choix ici sont tous dictés par le parsing, pas par l'esthétique :
    flux de texte unique (pas de colonne latérale à entrelacer), intitulés de
    section canoniques, coordonnées en clair sur les premières lignes, puces
    écrites « - » DANS la chaîne (une puce dessinée en cellule séparée ressort
    détachée de son texte à l'extraction), aucun aplat de couleur.

    `scale` resserre uniformément corps et interlignes — utilisé par `cv_pdf`
    pour tenter de faire tenir le CV sur une page (voir `_ats_autofit`)."""
    from fpdf.enums import XPos, YPos

    L = _ATS_LABELS[lg]
    W = _PAGE_W - 2 * _M
    pdf.set_auto_page_break(True, margin=_M)
    pdf.set_text_color(*_INK)

    def line(txt, size=9.0, style="", h=4.1, gap=0.0):
        txt = _latin(txt).strip()
        if not txt:
            return
        pdf.set_font("Helvetica", style, max(7.4, size * scale))
        pdf.set_x(_M)
        pdf.multi_cell(W, h * scale, txt, new_x=XPos.LMARGIN, new_y=YPos.NEXT)
        if gap:
            pdf.ln(gap * scale)

    def heading(txt):
        pdf.ln(2.0 * scale)
        line(txt, size=9.8, style="B", h=4.6)
        pdf.set_draw_color(*_LINE)
        pdf.line(_M, pdf.get_y() + 0.2, _M + W, pdf.get_y() + 0.2)
        pdf.ln(1.2 * scale)

    # ── En-tête : nom, titre, coordonnées en texte brut ──
    line(s.get("name") or "", size=16, style="B", h=7.4)
    line(s.get("role") or "", size=10.5, h=5.0)
    c = s.get("contact") or {}
    inline = _clean([c.get("email"), c.get("phone"), c.get("location")])
    if inline:
        line(" | ".join(str(x) for x in inline), size=9)
    for key in ("linkedin", "github", "portfolio"):     # un lien par ligne : jamais recollés
        if c.get(key):
            line(str(c[key]), size=9)

    if s.get("summary"):
        heading(L["profile"])
        line(str(s["summary"]))

    # Compétences juste après le profil : les ATS pondèrent ce qui apparaît tôt.
    skills = _clean(s.get("skills"))
    if skills:
        heading(L["skills"])
        for g in skills:
            if isinstance(g, dict):
                items = ", ".join(str(i) for i in _clean(g.get("items")))
                grp = str(g.get("group") or "").strip()
                line(f"{grp} : {items}" if grp and items else (items or grp))
            else:
                line(str(g))

    for key, label in (("experiences", L["experience"]), ("projects", L["projects"])):
        entries = _clean(s.get(key))
        if not entries:
            continue
        heading(label)
        for e in entries:
            if not isinstance(e, dict):
                line("- " + str(e))
                continue
            title = str(e.get("title") or "").strip()
            org = str(e.get("org") or e.get("meta") or "").strip()
            line(f"{title} - {org}" if org and title else (title or org), style="B", gap=0.3)
            meta = " | ".join(_clean([e.get("date"), e.get("stack")]))
            if meta:
                line(meta, size=8.8)
            for b in _clean(e.get("bullets"))[:8]:
                line("- " + str(b), h=3.95)
            pdf.ln(0.9)

    edu = _clean(s.get("education"))
    if edu:
        heading(L["education"])
        for e in edu:
            if isinstance(e, dict):
                line(str(e.get("title") or ""), style="B")
                if e.get("meta"):
                    line(str(e["meta"]), size=8.8)
            else:
                line(str(e))

    certs = _clean(s.get("certifications"))
    if certs:
        heading(L["certs"])
        for e in certs:
            line("- " + str(e), h=3.95)
    langs = _clean(s.get("languages"))
    if langs:
        # une seule ligne séparée par des virgules : aussi bien parsée qu'une liste
        # à puces, pour un tiers de la hauteur
        heading(L["langs"])
        line(", ".join(str(x) for x in langs), h=3.95)

    return bytes(pdf.output())


def _cv_pdf_fallback(pdf, cv_markdown: str) -> bytes:
    """Repli 1 colonne (Markdown brut) si le JSON structuré n'est pas disponible."""
    from fpdf.enums import XPos, YPos

    def cell(h, s):
        # new_x explicite : selon la version de fpdf2, multi_cell laisse sinon le
        # curseur en fin de ligne → « Not enough horizontal space » au bloc suivant
        pdf.multi_cell(0, h, s, new_x=XPos.LMARGIN, new_y=YPos.NEXT)

    pdf.set_auto_page_break(True, margin=_M)
    pdf.set_text_color(*_INK)
    for raw in (cv_markdown or "").splitlines():
        line = raw.rstrip()
        if not line.strip():
            pdf.ln(2.2)
            continue
        if line.startswith("# "):
            pdf.set_font("Helvetica", "B", 16)
            pdf.set_text_color(*_ACCENT)
            cell(7, _latin(line[2:]))
            pdf.set_text_color(*_INK)
        elif line.startswith("## "):
            pdf.ln(1.5)
            pdf.set_font("Helvetica", "B", 11.5)
            pdf.set_text_color(*_ACCENT)
            cell(5.5, _latin(line[3:].upper()))
            pdf.set_text_color(*_INK)
        elif line.lstrip().startswith(("- ", "* ")):
            pdf.set_font("Helvetica", "", 9.2)
            cell(4.4, _latin("- " + line.lstrip()[2:].replace("**", "")))
        else:
            pdf.set_font("Helvetica", "", 9.2)
            cell(4.4, _latin(line.replace("**", "")))
    return bytes(pdf.output())
